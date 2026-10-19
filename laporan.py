"""Laporan Word (.docx): hasil, perbandingan, uji galat tipe I (tabel + grafik), rangkuman proses, dan narasi
untuk menu Baseline (B1-B3), N1 SD-Hybrid, N2 TO-DAG, dan N3 MLC-DAG. Semua isi dibaca dari folder hasil/."""
import io
import re
import time
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
import config as C
import narasi as NR

BAGIAN = {"prep": "Persiapan Data", "eda": "EDA", "kausal": "Inferensi Kausal",
          "baseline": "Baseline (B1–B3)", "sd_hybrid": "N1 SD-Hybrid", "to_dag": "N2 TO-DAG",
          "mlc_dag": "N3 MLC-DAG", "sol": "SOL-DAG", "benchmark": "Benchmark semi-sintetik (H1–H5)",
          "benchmark_v4": "Benchmark v4 (M0–M6, H10–H13)"}
METODE_GRAF = ("baseline", "sd_hybrid", "to_dag", "mlc_dag")     # menu yang membandingkan metode penemuan graf
LEBAR = Cm(16)
NAMA_FILE = {"prep": "Laporan_Persiapan_Data", "eda": "Laporan_EDA", "kausal": "Laporan_Inferensi_Kausal",
             "baseline": "Laporan_Baseline_B1-B3", "sd_hybrid": "Laporan_N1_SD-Hybrid",
             "to_dag": "Laporan_N2_TO-DAG", "mlc_dag": "Laporan_N3_MLC-DAG",
             "sol": "Laporan_SOL-DAG", "benchmark": "Laporan_Benchmark_H1-H5",
             "benchmark_v4": "Laporan_Benchmark_v4_M0-M6",
             "lengkap": "Laporan_Lengkap_Semua_Menu"}
FOLDER = C.OUT / "laporan"


def berkas(kunci):
    """Path .docx untuk satu menu (kunci BAGIAN) atau 'lengkap'."""
    FOLDER.mkdir(parents=True, exist_ok=True)
    return FOLDER / f"{NAMA_FILE[kunci]}.docx"


def nama_sebagian(bagian, analisis):
    """Nama berkas laporan sebagian: gabungan nama menu (+ analisis bila tidak semua)."""
    s = "Laporan_" + "_".join(NAMA_FILE[b].removeprefix("Laporan_") for b in bagian)
    return s + ("" if set(analisis) == set(C.ANALISIS) else "_" + "_".join(analisis)) + ".docx"


def buat_menu(kunci):
    """Buat dan simpan laporan satu menu (atau 'lengkap') ke hasil/laporan/<nama menu>.docx."""
    return buat(tuple(BAGIAN) if kunci == "lengkap" else (kunci,), simpan=berkas(kunci))


# ================================================================== alat bantu dokumen
def _csv(p, **kw):
    try:
        return pd.read_csv(p, **kw) if p.exists() else None
    except pd.errors.EmptyDataError:
        return None


def _angka(v, fmt):
    if isinstance(v, (float, np.floating)) and np.isnan(v):
        return "–"
    if fmt is None:
        return str(v)
    try:
        t = fmt.format(v).replace(",", "#").replace(".", ",").replace("#", ".")
    except (ValueError, TypeError):
        t = str(v)
    return {"True": "ya", "False": "tidak"}.get(t, t.replace("-", "−") if re.match(r"^-\d", t) else t)


def gabung_ci(df):
    """RD, RD_lo, RD_hi -> satu kolom 'RD [CI 95%]' (idem RR) agar tabel tidak terlalu lebar."""
    df = df.copy()
    for k, f in (("RD", "{:+.4f}"), ("RR", "{:.2f}")):
        if {k, f"{k}_lo", f"{k}_hi"} <= set(df):
            df[f"{k} [CI 95%]"] = [f"{_angka(a, f)} [{_angka(b, f)}; {_angka(c, f)}]" if b == b else _angka(a, f)
                                   for a, b, c in zip(df[k], df[f"{k}_lo"], df[f"{k}_hi"])]
            pos = list(df).index(k)
            df = df.drop(columns=[k, f"{k}_lo", f"{k}_hi"])
            kol = list(df); kol.insert(pos, kol.pop(kol.index(f"{k} [CI 95%]"))); df = df[kol]
    return df


def _latar(sel, warna):
    tc = sel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), warna)
    tc.append(shd)


def tabel(doc, df, judul, fmt=None, sorot=None, ukuran=8):
    """df -> tabel Word bergaris; fmt: {kolom: '{:.3f}'}; sorot: Series bool baris yang diberi latar hijau."""
    if df is None or not len(df):
        return
    keterangan(doc, judul)
    fmt = fmt or {}
    ukuran = ukuran if len(df.columns) <= 8 else 7
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    for j, c in enumerate(df.columns):
        sel = t.rows[0].cells[j]; sel.text = ""
        r = sel.paragraphs[0].add_run(str(c).replace("_", " ")); r.bold = True; r.font.size = Pt(ukuran)
        _latar(sel, "D9E2F3")
    for i, (_, baris) in enumerate(df.iterrows()):
        cells = t.add_row().cells
        for j, c in enumerate(df.columns):
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(_angka(baris[c], fmt.get(c)))
            r.font.size = Pt(ukuran)
            if sorot is not None and bool(sorot.iloc[i]):
                _latar(cells[j], "D4EDDA")
    doc.add_paragraph()


def keterangan(doc, teks):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    r = p.add_run(teks); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)


def gambar(doc, sumber, judul, lebar=LEBAR):
    """sumber: path PNG atau figure matplotlib."""
    if sumber is None:
        return
    if hasattr(sumber, "savefig"):
        buf = io.BytesIO(); sumber.savefig(buf, dpi=150, bbox_inches="tight"); plt.close(sumber); buf.seek(0)
        sumber = buf
    elif not sumber.exists():
        return
    else:
        sumber = str(sumber)
    doc.add_picture(sumber, width=lebar)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(judul); r.italic = True; r.font.size = Pt(8)


def markdown(par, teks):
    """**tebal** dan *miring* sederhana -> run Word."""
    for bag in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", teks):
        if bag.startswith("**") and bag.endswith("**"):
            par.add_run(bag[2:-2]).bold = True
        elif bag.startswith("*") and bag.endswith("*") and len(bag) > 1:
            par.add_run(bag[1:-1]).italic = True
        elif bag:
            par.add_run(bag)


def narasi(doc, r):
    if r is None:
        doc.add_paragraph("Hasil belum tersedia; jalankan menu ini di aplikasi terlebih dahulu.")
        return
    doc.add_heading("Rangkuman proses", level=3)
    for i, s in enumerate(r['proses'], 1):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6); p.paragraph_format.space_after = Pt(2)
        markdown(p, f"{i}.	{s}")
    doc.add_heading("Narasi", level=3)
    for s in r['narasi']:
        p = doc.add_paragraph(); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        markdown(p, s)


# ================================================================== grafik tambahan
def fig_tipe1(sim):
    fig, ax = plt.subplots(figsize=(8, 3.2))
    y = np.arange(len(sim))
    ax.barh(y + 0.2, sim['naif (B2)'] * 100, 0.4, color="#d62728", label="naif (B2)")
    ax.barh(y - 0.2, sim['terkoreksi desain (N1)'] * 100, 0.4, color="#1f77b4", label="terkoreksi desain (N1)")
    ax.axvline(sim.nominal.iloc[0] * 100, color="#333333", ls="--", lw=1, label="nominal")
    ax.set_yticks(y); ax.set_yticklabels(sim.uji, fontsize=8); ax.set_xlabel("laju galat tipe I (%)")
    ax.legend(fontsize=8); fig.tight_layout()
    return fig


def fig_barh(s, xlabel, warna="#1f77b4"):
    s = s.sort_values()
    fig, ax = plt.subplots(figsize=(7, 0.28 * len(s) + 1))
    ax.barh(s.index, s.values, color=warna); ax.set_xlabel(xlabel); ax.tick_params(labelsize=8)
    fig.tight_layout()
    return fig


def fig_heatmap(df, judul):
    fig, ax = plt.subplots(figsize=(6.5, 5.5))
    im = ax.imshow(df.values, cmap="RdBu_r", vmin=-1, vmax=1)
    ax.set_xticks(range(len(df))); ax.set_xticklabels(df.columns, rotation=90, fontsize=7)
    ax.set_yticks(range(len(df))); ax.set_yticklabels(df.index, fontsize=7)
    for i in range(len(df)):
        for j in range(len(df)):
            ax.text(j, i, f"{df.values[i, j]:.2f}", ha="center", va="center", fontsize=5)
    fig.colorbar(im, fraction=0.04); ax.set_title(judul, fontsize=9); fig.tight_layout()
    return fig


def fig_garis(x, y, xlabel, ylabel):
    fig, ax = plt.subplots(figsize=(7, 2.8))
    ax.plot(x, y, marker="o", ms=3); ax.set_xlabel(xlabel); ax.set_ylabel(ylabel); fig.tight_layout()
    return fig


# ================================================================== isi per menu
F_RD = {"RD": "{:+.4f}", "RD_lo": "{:+.4f}", "RD_hi": "{:+.4f}", "RR": "{:.2f}", "RR_lo": "{:.2f}", "RR_hi": "{:.2f}",
        "lebar_CI": "{:.4f}", "BIC_SEM": "{:.1f}", "CFI": "{:.3f}", "RMSEA": "{:.3f}", "SRMR": "{:.3f}",
        "detik_penemuan": "{:.3f}", "memori_MB": "{:.2f}", "peringkat_rata2": "{:.2f}", "sisi": "{:.0f}",
        "SHD": "{:.0f}", "node": "{:.0f}", "estimasi": "{:+.4f}", "se_boot": "{:.4f}", "ci_bawah": "{:+.4f}",
        "ci_atas": "{:+.4f}", "beta": "{:+.4f}", "beta_std": "{:+.3f}", "muatan": "{:.3f}", "galat": "{:.3f}",
        "tetrad": "{:.3f}", "tetrad_lintas": "{:.3f}", "p1": "{:.4f}", "p0": "{:.4f}", "F_tier": "{:.1f}",
        "naif (B2)": "{:.1%}", "terkoreksi desain (N1)": "{:.1%}", "nominal": "{:.0%}", "deff": "{:.2f}",
        **{f"rank_{k}": "{:.0f}" for k in ("BIC_SEM", "CFI", "lebar_CI", "detik_penemuan", "memori_MB")}}


def kv_tabel(doc, path, kunci, judul):
    s = _csv(path, index_col=0)
    if s is None:
        return
    s = s.nilai
    df = pd.DataFrame({"ukuran": [k.replace("_", " ") for k in kunci if k in s.index],
                       "nilai": [re.sub(r"(\d)\.(\d)", r"\1,\2", str(s[k])) for k in kunci if k in s.index]})
    tabel(doc, df, judul)


def isi_baseline(doc, n, nomor):
    out = C.OUT / "baseline" / n
    narasi(doc, NR.baseline(n))
    rk = _csv(out / "rangkuman.csv")
    if rk is None:
        return
    doc.add_heading("Hasil dan perbandingan", level=3)
    kol = [c for c in ["metode", "sisi", "SHD", "BIC_SEM", "CFI", "RMSEA", "RD", "RD_lo", "RD_hi", "lebar_CI",
                       "detik_penemuan", "peringkat_rata2"] if c in rk]
    tabel(doc, gabung_ci(rk[kol]), f"Tabel {nomor}.1 Pemeringkatan B1–B3 (hijau = terbaik; BIC-SEM ↓, CFI ↑, lebar CI ↓)",
          F_RD, rk.terbaik)
    gambar(doc, out / "perbandingan.png", f"Gambar {nomor}.1 Risk difference dan CI 95% bootstrap PSU per metode")
    fit = _csv(out / "b1_kecocokan.csv")
    if fit is not None:
        tabel(doc, fit.rename(columns={"ukuran": "ukuran kecocokan"}),
              f"Tabel {nomor}.2 Kecocokan DAG pakar (B1 SEM)", {"nilai": "{:.3f}"})
    tabel(doc, _csv(out / "b1_dekomposisi.csv"), f"Tabel {nomor}.3 Dekomposisi efek B1 (linear dan probit)", F_RD)
    ida = _csv(out / "b2_ida.csv")
    tabel(doc, None if ida is None else gabung_ci(ida), f"Tabel {nomor}.4 Himpunan efek IDA (B2)", F_RD)
    kv_tabel(doc, out / "b3_info.csv", ["skor", "awal", "langkah", "skor_akhir", "skor_DAG_pakar", "sisi", "SHD",
                                        "adjustment", "catatan"], f"Tabel {nomor}.5 Ringkasan B3 (BN berbasis skor)")
    for k, (f, j) in enumerate([("b1_graf.png", "DAG pakar (B1)"), ("b2_graf.png", "CPDAG hasil PC + tier (B2)"),
                                ("b3_graf.png", "DAG hasil Hill-Climbing (B3)")], start=2):
        gambar(doc, out / f, f"Gambar {nomor}.{k} {j}")


def banding_tabel(doc, out, nomor, kriteria):
    bd = _csv(out / "banding.csv")
    if bd is None:
        return
    utama = {"BIC_SEM": ["BIC_SEM", "CFI"], "detik_penemuan": ["detik_penemuan", "memori_MB"]}
    kol = ["metode", "node" if "node" in bd and bd.node.notna().any() else "sisi"]
    for k in kriteria:
        kol += [c for c in utama.get(k, [k]) if c in bd and c not in kol and c != "lebar_CI"]
    kol += ["RD", "RD_lo", "RD_hi", "lebar_CI", "peringkat_rata2"]
    tabel(doc, gabung_ci(bd[[c for c in kol if c in bd]]), f"Tabel {nomor}.1 Perbandingan dengan baseline B1–B3 "
          "(hijau = terbaik; kriteria: " + ", ".join(NR.NK.get(k, k) for k in kriteria)
          + "; B1 tidak diperingkat pada waktu dan memori)", F_RD, bd.terbaik)
    gambar(doc, out / "banding.png", f"Gambar {nomor}.1 Risk difference dan CI 95%: kandidat vs baseline")


def isi_novelty(doc, metode, n, nomor):
    import novelty
    out = C.OUT / "novelty" / metode / n
    pre = {"sd_hybrid": "sd", "to_dag": "to", "mlc_dag": "mlc"}[metode]
    narasi(doc, NR.novelty(metode, n))
    if not (out / f"{pre}_info.csv").exists():
        return
    doc.add_heading("Hasil dan perbandingan", level=3)
    banding_tabel(doc, out, nomor, list(novelty.METODE[metode]['kriteria']))
    if metode == "sd_hybrid":
        sim = _csv(out / "sd_tipe1.csv")
        if sim is not None:
            doc.add_heading("Uji galat tipe I", level=3)
            tabel(doc, sim, f"Tabel {nomor}.2 Laju galat tipe I di bawah H0 (desain IDHS nyata, "
                  f"{C.NOVELTY['sim_rep']} replikasi, ICC {str(C.NOVELTY['sim_icc']).replace('.', ',')})", F_RD)
            gambar(doc, fig_tipe1(sim), f"Gambar {nomor}.2 Galat tipe I: uji naif vs terkoreksi desain")
        kv_tabel(doc, out / "sd_info.csv", ["jumlah_uji", "sisi_skeleton", "sisi_skeleton_B2_naif", "delta_bar",
                                            "n_Kish", "n_eff", "df_desain", "skor", "sisi", "SHD", "adjustment",
                                            "catatan", "detik_penemuan", "memori_MB"],
                 f"Tabel {nomor}.3 Ringkasan SD-Hybrid")
        dv = _csv(out / "sd_deff.csv", index_col=0)
        if dv is not None:
            gambar(doc, fig_barh(dv.deff, "design effect klaster"), f"Gambar {nomor}.3 Design effect per variabel")
        gambar(doc, out / "sd_graf.png", f"Gambar {nomor}.4 DAG hasil SD-Hybrid")
    elif metode == "to_dag":
        kv_tabel(doc, out / "to_info.csv", ["ukuran_tier", "ruang_DAG", "ruang_urutan_tanpa_tier", "ruang_urutan_tier",
                                            "ruang_setelah_dekomposisi", "skor_lokal_dihitung", "grow_shrink_dihitung",
                                            "skor_TO_DAG", "skor_B3_pada_skor_sama", "skor_pakar_pada_skor_sama",
                                            "sisi", "SHD", "adjustment", "catatan", "detik_penemuan", "memori_MB"],
                 f"Tabel {nomor}.2 Ruang pencarian, skor, dan sumber daya TO-DAG")
        tabel(doc, _csv(out / "to_jejak.csv"), f"Tabel {nomor}.3 Jejak konvergensi per tier", F_RD)
        tabel(doc, _csv(out / "to_urutan.csv"), f"Tabel {nomor}.4 Urutan topologis terpilih")
        gambar(doc, out / "to_graf.png", f"Gambar {nomor}.2 DAG hasil TO-DAG")
    else:
        kv_tabel(doc, out / "mlc_info.csv", ["node_tanpa_konstruk", "node_struktural", "indikator_awal",
                                             "indikator_dibuang", "reliabilitas_omega", "konstruk_bermasalah",
                                             "detik_korelasi", "detik_model_ukur", "detik_struktur",
                                             "detik_struktur_tanpa_konstruk", "uji_atau_skor",
                                             "uji_atau_skor_tanpa_konstruk", "sisi", "adjustment", "catatan"],
                 f"Tabel {nomor}.2 Ringkasan MLC-DAG dan ablasi tanpa konstruk")
        tabel(doc, _csv(out / "mlc_muatan.csv"), f"Tabel {nomor}.3 Muatan faktor setelah purifikasi", F_RD)
        pl = _csv(out / "mlc_purifikasi.csv")
        if pl is not None and len(pl):
            tabel(doc, pl, f"Tabel {nomor}.4 Indikator yang dibuang saat purifikasi", F_RD)
        sg = _csv(out / "mlc_sigma.csv", index_col=0)
        if sg is not None:
            gambar(doc, fig_heatmap(sg, "Σ̂_V"), f"Gambar {nomor}.2 Korelasi konstruk (terkoreksi galat ukur) dan "
                                                 "variabel teramati", Cm(13))
        gambar(doc, out / "mlc_graf.png", f"Gambar {nomor}.3 DAG struktural hasil MLC-DAG")


# ================================================================== menu 1-3
def _data():
    p = C.CACHE / "analisis.pkl"
    return pd.read_pickle(p) if p.exists() else None


def isi_prep(doc, nomor):
    df = _data()
    narasi(doc, NR.prep(df))
    if df is None:
        return
    doc.add_heading("Hasil", level=3)
    ringkas = pd.DataFrame({"ukuran": ["kelahiran (5 tahun terakhir)", "variabel", "klaster", "PSU", "strata",
                                       "provinsi", "jumlah bobot sampel"],
                            "nilai": [len(df), df.shape[1], df.klaster.nunique(), df.groupby(['strata', 'psu']).ngroups,
                                      df.strata.nunique(), df.kode_prov_v024.nunique(), round(df.bobot.sum(), 1)]})
    tabel(doc, ringkas, f"Tabel {nomor}.1 Ringkasan dataset analisis", {"nilai": "{:,.0f}"})
    kunci = [c for c in df if c.startswith(('T_', 'Y_', 'M_'))]
    w = df.bobot
    tt = pd.DataFrame({"variabel": kunci, "label": [NR.lbl(c) for c in kunci],
                       "peran": ["treatment" if c.startswith('T_') else "outcome" if c.startswith('Y_') else "mediator"
                                 for c in kunci],
                       "n terisi": [int(df[c].notna().sum()) for c in kunci],
                       "kosong": [df[c].isna().mean() for c in kunci],
                       "rerata berbobot": [np.average(df[c].dropna(), weights=w[df[c].notna()]) for c in kunci]})
    tabel(doc, tt, f"Tabel {nomor}.2 Treatment, outcome, dan mediator", {"n terisi": "{:,.0f}", "kosong": "{:.1%}",
                                                                        "rerata berbobot": "{:.3f}"})
    ks = df.isna().mean().sort_values(ascending=False)
    ks = ks[ks > 0].head(20)
    if len(ks):
        gambar(doc, fig_barh(ks * 100, "% data kosong", "#d62728"), f"Gambar {nomor}.1 Variabel dengan data kosong "
                                                                     "terbanyak (maksimum 20)")
    kamus = pd.DataFrame({"variabel": df.columns, "tipe": df.dtypes.astype(str).values,
                          "n terisi": df.notna().sum().values, "unik": df.nunique().values,
                          "kelompok": ["Treatment" if k.startswith("T_") else "Outcome" if k.startswith("Y_")
                                       else "Mediator" if k.startswith("M_") else "Indikator MLC" if k.startswith("ind_")
                                       else "Kovariat" for k in df.columns]})
    tabel(doc, kamus, f"Tabel {nomor}.3 Kamus variabel", {"n terisi": "{:,.0f}", "unik": "{:,.0f}"}, ukuran=7)


def isi_eda(doc, nomor):
    o = C.OUT / "eda"
    narasi(doc, NR.eda())
    desk = _csv(o / "02_deskriptif_svy.csv")
    if desk is None:
        return
    doc.add_heading("Hasil", level=3)
    tabel(doc, desk.assign(label=desk.variabel.map(NR.lbl))[["variabel", "label", "rerata", "se", "ci_bawah", "ci_atas",
                                                            "n"]],
          f"Tabel {nomor}.1 Rerata berbobot dan SE desain (linearisasi Taylor, strata-PSU)",
          {"rerata": "{:.3f}", "se": "{:.4f}", "ci_bawah": "{:.3f}", "ci_atas": "{:.3f}", "n": "{:,.0f}"})
    q = _csv(o / "03_per_kuintil.csv")
    if q is not None:
        tabel(doc, q, f"Tabel {nomor}.2 Treatment dan outcome per kuintil kekayaan non-WASH (proporsi berbobot)",
              {q.columns[0]: "{:.0f}", **{c: "{:.3f}" for c in q.columns[1:]}})
    gambar(doc, o / "06_kuintil.png", f"Gambar {nomor}.1 Treatment dan diare per kuintil kekayaan non-WASH")
    pv = _csv(o / "04_per_provinsi.csv")
    if pv is not None:
        tabel(doc, pv, f"Tabel {nomor}.3 Treatment dan outcome per provinsi (proporsi berbobot)",
              {c: "{:.3f}" for c in pv.columns[1:]}, ukuran=7)
        ts = [t for t in ("T_sanitasi", "T_pbi", "T_fasilitas") if t in pv]
        pv2 = pv.set_index(pv.columns[0])[ts].sort_values(ts[0])
        fig, ax = plt.subplots(figsize=(7.5, 9))
        y = np.arange(len(pv2)); h = 0.8 / len(ts)
        for k, (t, c) in enumerate(zip(ts, ["#1f77b4", "#ff7f0e", "#2ca02c"])):
            ax.barh(y + (k - 1) * h, pv2[t] * 100, h, color=c, label=NR.lbl(t))
        ax.set_yticks(y); ax.set_yticklabels(pv2.index, fontsize=7); ax.set_xlabel("% berbobot")
        ax.legend(fontsize=8, loc="lower right"); fig.tight_layout()
        gambar(doc, fig, f"Gambar {nomor}.2 Cakupan treatment per provinsi", Cm(14))
    bal = _csv(o / "05_smd_per_treatment.csv")
    if bal is not None:
        bal = bal.rename(columns={bal.columns[0]: "kovariat"})
        tabel(doc, bal, f"Tabel {nomor}.4 Standardized mean difference (SMD) berbobot kovariat per treatment "
              "(|SMD| > 0,1 = tidak seimbang)", {c: "{:+.3f}" for c in bal.columns[1:]}, ukuran=7)
    gambar(doc, o / "07_korelasi.png", f"Gambar {nomor}.3 Korelasi antar variabel utama", Cm(12))
    ms = _csv(o / "01_missing.csv")
    if ms is not None:
        ms = ms.rename(columns={ms.columns[0]: "variabel"})
        tabel(doc, ms[ms.prop_kosong > 0].head(20), f"Tabel {nomor}.5 Proporsi data kosong (20 terbesar)",
              {"prop_kosong": "{:.1%}"})


def isi_kausal(doc, n, nomor):
    out = C.OUT / "kausal" / n
    narasi(doc, NR.kausal(n))
    adj = NR._adj(out / "adjustment.txt")
    if adj is None:
        return
    doc.add_heading("Hasil", level=3)
    tabel(doc, pd.DataFrame({"komponen": list(adj), "nilai": list(adj.values())}),
          f"Tabel {nomor}.1 Identifikasi dan adjustment set (kriteria backdoor pada DAG pakar)")
    gambar(doc, out / "dag.png", f"Gambar {nomor}.1 DAG pakar: merah = treatment, hijau = outcome, biru = dikontrol, "
                                 "oranye = tidak dikontrol, abu = laten")
    est = _csv(out / "estimasi.csv", index_col=0)
    if est is None:
        return
    rows = []
    for k, nm in (("gcomp", "g-komputasi"), ("ipw", "IPW"), ("aipw", "AIPW (utama)")):
        g = lambda s, c: est.loc[f"{k}_{s}", c]
        rows.append({"estimator": nm, "p1": g("p1", "estimasi"), "p0": g("p0", "estimasi"),
                     "RD": g("RD", "estimasi"), "RD_lo": g("RD", "ci_bawah"), "RD_hi": g("RD", "ci_atas"),
                     "RR": g("RR", "estimasi"), "RR_lo": g("RR", "ci_bawah"), "RR_hi": g("RR", "ci_atas"),
                     "se_boot": g("RD", "se_boot")})
    tabel(doc, gabung_ci(pd.DataFrame(rows)), f"Tabel {nomor}.2 Estimasi efek (CI 95% bootstrap PSU, "
          f"{C.BOOT} replikasi)", F_RD)
    rr, lo, hi = est.loc["aipw_RR", "estimasi"], est.loc["aipw_RR", "ci_bawah"], est.loc["aipw_RR", "ci_atas"]
    ev_ci = 1.0 if lo <= 1 <= hi else NR.evalue(lo if rr > 1 else hi)
    tabel(doc, pd.DataFrame({"ukuran": ["E-value estimasi titik", "E-value batas CI"],
                             "nilai": [NR.evalue(rr), ev_ci]}),
          f"Tabel {nomor}.3 Sensitivitas terhadap confounder tak terukur (E-value, dari RR AIPW)", {"nilai": "{:.2f}"})
    bal = _csv(out / "keseimbangan.csv", index_col=0)
    if bal is not None:
        tabel(doc, bal.reset_index().rename(columns={"index": "kovariat"}),
              f"Tabel {nomor}.4 Keseimbangan kovariat sebelum dan sesudah IPW (SMD; target |SMD| < 0,1)",
              {"smd_awal": "{:+.3f}", "smd_ipw": "{:+.3f}"})
        b = bal.abs().sort_values("smd_awal")
        fig, ax = plt.subplots(figsize=(7, 0.3 * len(b) + 1))
        y = np.arange(len(b))
        ax.barh(y + 0.2, b.smd_awal, 0.4, color="#d62728", label="sebelum")
        ax.barh(y - 0.2, b.smd_ipw, 0.4, color="#1f77b4", label="sesudah IPW")
        ax.axvline(0.1, color="#333333", ls="--", lw=1)
        ax.set_yticks(y); ax.set_yticklabels(b.index, fontsize=7); ax.set_xlabel("|SMD|"); ax.legend(fontsize=8)
        fig.tight_layout()
        gambar(doc, fig, f"Gambar {nomor}.2 |SMD| kovariat sebelum dan sesudah pembobotan IPW")
    gambar(doc, out / "overlap.png", f"Gambar {nomor}.3 Overlap propensity score per kelompok treatment", Cm(12))


def isi_sol(doc, n, nomor):
    out = C.OUT / "sol_dag" / n
    narasi(doc, NR.sol(n))
    v = _csv(out / "sol_varian.csv")
    if v is None:
        return
    doc.add_heading("Hasil", level=3)
    kol = ["varian", "node", "sisi", "sisi_skeleton", "jumlah_uji", "n_eff", "RD", "RD_lo", "RD_hi", "detik_penemuan",
           "memori_MB"]
    tabel(doc, gabung_ci(v[kol]), f"Tabel {nomor}.1 SOL-DAG dan ablasinya", {**F_RD, "n_eff": "{:,.0f}",
                                                                          "node": "{:.0f}", "sisi_skeleton": "{:.0f}",
                                                                          "jumlah_uji": "{:,.0f}"})
    bd = _csv(out / "sol_banding.csv")
    if bd is not None:
        tabel(doc, gabung_ci(bd.drop(columns=[c for c in ("kode",) if c in bd])),
              f"Tabel {nomor}.2 SOL-DAG dibanding baseline dan kandidat (data nyata)", F_RD)
        gambar(doc, out / "sol_banding.png", f"Gambar {nomor}.1 Risk difference dan CI 95% per metode")
    gambar(doc, out / "sol_graf.png", f"Gambar {nomor}.2 DAG hasil SOL-DAG")
    v2 = _csv(out / "sol_v2.csv")
    if v2 is not None:
        doc.add_heading("SOL v2 (pra-registrasi)", level=3)
        tabel(doc, gabung_ci(v2[["varian", "node", "sisi", "tak_berarah", "RD", "RD_lo", "RD_hi", "himpunan_kandidat",
                                 "induk_T", "detik"]]), f"Tabel {nomor}.3 SOL v2: efek selalu diestimasi (AIPW cross-fitting)",
              {**F_RD, "node": "{:.0f}", "tak_berarah": "{:.0f}", "himpunan_kandidat": "{:.0f}"})
        gambar(doc, out / "sol_v2_graf.png", f"Gambar {nomor}.3 Graf SOL v2 (uji-t desain)")


def isi_benchmark(doc, n, nomor):
    import json
    out = C.OUT / "benchmark" / n
    narasi(doc, NR.benchmark(n))
    pt = _csv(out / "putusan.csv")
    if pt is None:
        return
    doc.add_heading("Putusan hipotesis", level=3)
    tabel(doc, pt, f"Tabel {nomor}.1 Putusan H1–H5 (SOL-DAG vs baseline)", sorot=pt.putusan == "didukung")
    rk = _csv(out / "ringkasan.csv")
    hp = _csv(out / "hipotesis.csv")
    no = 2
    for s in ("A", "B"):
        if rk is None or not (rk.skenario == s).any():
            continue
        doc.add_heading(f"Skenario {s}", level=3)
        r = rk[rk.skenario == s]
        kol = ["metode", "SHD", "SID", "RMSE", "bias", "efek_nol", "cakupan", "detik", "memori", "stabilitas_SHD"]
        tabel(doc, r[kol], f"Tabel {nomor}.{no} Metrik rata-rata per metode, skenario {s}",
              {"SHD": "{:.1f}", "SID": "{:.1f}", "RMSE": "{:.4f}", "bias": "{:+.4f}", "efek_nol": "{:.0%}",
               "cakupan": "{:.0%}", "detik": "{:.2f}", "memori": "{:.1f}", "stabilitas_SHD": "{:.1f}"}); no += 1
        fig, ax = plt.subplots(1, 2, figsize=(10, 0.35 * len(r) + 1.5))
        rr = r.set_index("metode")
        ax[0].barh(rr.index, rr.SHD, color="#1f77b4"); ax[0].set_xlabel("SHD rata-rata (↓)")
        ax[1].barh(rr.index, rr.RMSE, color="#ff7f0e"); ax[1].set_xlabel("RMSE efek (↓)"); ax[1].set_yticklabels([])
        for a_ in ax:
            a_.tick_params(labelsize=8); a_.invert_yaxis()
        fig.tight_layout()
        gambar(doc, fig, f"Gambar {nomor}.{no} Akurasi struktur dan efek, skenario {s}")
        if hp is not None and (hp.skenario == s).any():
            h = hp[hp.skenario == s][["hipotesis", "pembanding", "metrik", "SOL", "lawan", "median_selisih", "p_holm",
                                      "keputusan"]]
            tabel(doc, h, f"Tabel {nomor}.{no} Uji berpasangan Wilcoxon satu sisi (SOL-DAG < pembanding), skenario {s}",
                  {"SOL": "{:.4g}", "lawan": "{:.4g}", "median_selisih": "{:+.4g}", "p_holm": "{:.4f}"},
                  sorot=h.keputusan == "SOL lebih baik", ukuran=7); no += 1
        gambar(doc, out / f"graf_benar_{s}.png", f"Gambar {nomor}.{no} Graf benar skenario {s}"); no += 1
    h5 = _csv(out / "hipotesis_H5.csv")
    if h5 is not None and len(h5):
        doc.add_heading("H5 — validitas koreksi desain", level=3)
        tabel(doc, h5, f"Tabel {nomor}.{no} Galat tipe I, daya uji, dan cakupan CI 95% (interval Wilson)",
              {"nilai": "{:.1%}", "ci_lo": "{:.1%}", "ci_hi": "{:.1%}", "target": "{:.0%}", "n": "{:,.0f}"},
              sorot=h5.sesuai_target.fillna(False).astype(bool))
        fig, ax = plt.subplots(figsize=(8, 0.35 * len(h5) + 1))
        lab = h5.ukuran + " · " + h5.uji
        ax.errorbar(h5.nilai * 100, range(len(h5)), xerr=[(h5.nilai - h5.ci_lo) * 100, (h5.ci_hi - h5.nilai) * 100],
                    fmt="o", capsize=3)
        for t in h5.target.dropna().unique():
            ax.axvline(t * 100, ls="--", color="#888888", lw=1)
        ax.set_yticks(range(len(h5))); ax.set_yticklabels(lab, fontsize=7); ax.invert_yaxis(); ax.set_xlabel("%")
        fig.tight_layout()
        gambar(doc, fig, f"Gambar {nomor}.{no} H5: proporsi dan interval Wilson terhadap target (garis putus-putus)")
    import benchmark as BM
    sel, hv = BM.evaluasi_v2()
    if len(hv):
        doc.add_heading("Benchmark v2 (pra-registrasi)", level=3)
        tabel(doc, hv, f"Tabel {nomor}.v2a Putusan H1v2–H6v2 menurut kriteria PRAREGISTRASI_v2.md",
              sorot=hv.putusan == "diterima")
        tabel(doc, sel, f"Tabel {nomor}.v2b Metrik per sel grid (ICC × β_TY) dan varian",
              {"SHD": "{:.2f}", "RMSE": "{:.4f}", "cakupan": "{:.1%}", "efek_nol": "{:.1%}", "detik": "{:.2f}",
               "MB": "{:.1f}", "t1_naif": "{:.1%}", "t1_desain": "{:.1%}", "daya_naif": "{:.1%}", "daya_desain": "{:.1%}",
               "icc": "{:.2f}", "beta_ty": "{:+.2f}", "rep": "{:.0f}"}, ukuran=7)
    kal3, hil3, h3 = BM.evaluasi_v3()
    if len(h3):
        doc.add_heading("Benchmark v3 (pra-registrasi S6–S10)", level=3)
        tabel(doc, h3, f"Tabel {nomor}.v3a Putusan H7a–H9 menurut kriteria PRAREGISTRASI_v3.md",
              sorot=h3.putusan.str.startswith("diterima"))
        doc.add_paragraph(BM.TAFSIR_V3).runs[0].italic = True
        if len(kal3):
            tabel(doc, kal3, f"Tabel {nomor}.v3b S8: galat tipe I dan daya lima uji per skenario",
                  {"tipeI": "{:.1%}", "wilson_lo": "{:.1%}", "wilson_hi": "{:.1%}", "daya": "{:.1%}", "n_H0": "{:.0f}",
                   "n_H1": "{:.0f}"})
            fig, ax = plt.subplots(figsize=(8, 3.6))
            for k_, (sc, g) in enumerate(kal3.groupby("skenario")):
                y = np.arange(len(g)) + k_ * 0.35
                ax.barh(y, g.tipeI * 100, 0.35, label=sc)
            ax.set_yticks(np.arange(len(g)) + 0.17); ax.set_yticklabels(g.uji, fontsize=8)
            ax.axvline(5, ls="--", color="#555555", lw=1); ax.set_xlabel("galat tipe I (%)"); ax.legend(fontsize=8)
            fig.tight_layout()
            gambar(doc, fig, f"Gambar {nomor}.v3 Galat tipe I lima uji (garis = nominal 5%)")
        if len(hil3):
            tabel(doc, hil3, f"Tabel {nomor}.v3c Dampak hilir (p = 15)", {c: "{:.3f}" for c in hil3.columns[2:]}, ukuran=7)


def isi_benchmark_v4(doc, nomor):
    """Benchmark v4 (bench_v4.py) + M0-M6 pada IDHS (idhs_m0m6.py); tidak per analisis."""
    out = C.OUT / "benchmark_v4"
    narasi(doc, NR.benchmark_v4())
    d, hp = NR.data_v4(), _csv(out / "hipotesis.csv")
    if d is None or hp is None:
        return
    F = {"valid": "{:.1%}", "valid_lo": "{:.1%}", "valid_hi": "{:.1%}", "bias": "{:+.4f}", "RMSE": "{:.4f}",
         "cakupan": "{:.1%}", "lebar_CI": "{:.4f}", "n_Z": "{:.1f}", "FDR": "{:.1%}", "recall": "{:.1%}", "replikasi": "{:.0f}"}
    doc.add_heading("Putusan hipotesis", level=3)
    tabel(doc, hp, f"Tabel {nomor}.1 Putusan H10–H13 menurut RENCANA_ANALISIS_v4.md (sel ICC heterogen)",
          sorot=hp.putusan == "diterima")
    doc.add_heading("Kinerja per metode", level=3)
    het = NR.kinerja_v4(d[d.icc == "heterogen"])
    tabel(doc, het, f"Tabel {nomor}.2 Kinerja per metode, ICC heterogen (interval Wilson untuk validitas)", F, ukuran=7)
    tabel(doc, NR.kinerja_v4(d[d.icc != "heterogen"]), f"Tabel {nomor}.3 Kinerja per metode, ICC 0,15", F, ukuran=7)
    h = het.set_index("metode")
    fig, ax = plt.subplots(1, 2, figsize=(10, 0.35 * len(h) + 1.5), sharey=True)
    y = np.arange(len(h))                                              # posisi numerik: baris FDR sejajar dengan metodenya
    ax[0].barh(y, h.valid * 100, color="#1f77b4",
               xerr=[(h.valid - h.valid_lo) * 100, (h.valid_hi - h.valid) * 100], ecolor="#333333", capsize=2)
    ax[0].axvline(90, ls="--", color="#555555", lw=1); ax[0].set_xlabel("himpunan penyesuaian valid, % (↑; garis = target H10)")
    ax[1].barh(y, h.FDR.fillna(0) * 100, color="#d62728"); ax[1].set_xlabel("FDR skeleton, % (↓)")
    ax[0].set_yticks(y); ax[0].set_yticklabels(h.index); ax[0].invert_yaxis()
    for a_ in ax:
        a_.tick_params(labelsize=8)
    fig.tight_layout()
    gambar(doc, fig, f"Gambar {nomor}.1 Validitas himpunan penyesuaian dan FDR skeleton, ICC heterogen (M0, M1, M5, M6 "
           "tidak menjalankan penemuan graf sendiri)")
    fig, ax = plt.subplots(1, 2, figsize=(10, 0.35 * len(h) + 1.5))
    ax[0].barh(h.index, h.RMSE, color="#ff7f0e"); ax[0].set_xlabel("RMSE risk difference (↓)")
    ax[1].barh(h.index, h.cakupan * 100, color="#2ca02c"); ax[1].axvline(95, ls="--", color="#555555", lw=1)
    ax[1].set_xlabel("cakupan CI 95%, % (garis = nominal)"); ax[1].set_yticklabels([])
    for a_ in ax:
        a_.tick_params(labelsize=8); a_.invert_yaxis()
    fig.tight_layout()
    gambar(doc, fig, f"Gambar {nomor}.2 RMSE dan cakupan CI 95% per metode, ICC heterogen")
    doc.add_heading("Validitas menurut faktor simulasi", level=3)
    urut = [m for m in NR.URUT_V4 if m in set(d.metode)]
    faktor = {"kuat": "kekuatan confounder", "p": "jumlah node", "beta_ty": "efek T→Y (laten probit)", "icc": "ICC"}
    for no, (f, judul) in enumerate(faktor.items(), start=4):
        pv = d.pivot_table(index="metode", columns=f, values="valid").reindex(urut)
        pv.columns = [str(c) for c in pv.columns]
        tabel(doc, pv.reset_index(), f"Tabel {nomor}.{no} Himpunan valid menurut {judul}, semua sel",
              {c: "{:.1%}" for c in pv.columns})
    k = d.pivot_table(index="metode", columns="kuat", values="valid").reindex(urut) * 100
    fig, ax = plt.subplots(figsize=(8, 0.4 * len(k) + 1.3)); y = np.arange(len(k))
    ax.barh(y - 0.2, k.kuat, 0.4, color="#1f77b4", label="confounder kuat")
    ax.barh(y + 0.2, k.lemah, 0.4, color="#aec7e8", label="confounder lemah")
    ax.set_yticks(y); ax.set_yticklabels(k.index, fontsize=8); ax.invert_yaxis()
    ax.set_xlabel("himpunan penyesuaian valid, %"); ax.legend(fontsize=8); fig.tight_layout()
    gambar(doc, fig, f"Gambar {nomor}.3 Validitas menurut kekuatan confounder, semua sel")
    no, ng = 8, 4
    for an in ("sanitasi_diare", "pbi_diare"):
        rd = _csv(C.OUT / "idhs_m0m6" / an / "rd.csv")
        if rd is None:
            continue
        if no == 8:
            doc.add_heading("M0–M6 pada IDHS 2017", level=3)
        tabel(doc, gabung_ci(rd[['metode', 'n_Z', 'RD', 'RD_lo', 'RD_hi', 'Z']]),
              f"Tabel {nomor}.{no} {an}: risk difference per strategi (AIPW cross-fitting per PSU, n = {NR.ribu(rd.n.iloc[0])})",
              {"n_Z": "{:.0f}"}, ukuran=7); no += 1
        gambar(doc, C.OUT / "idhs_m0m6" / an / "forest.png", f"Gambar {nomor}.{ng} {an}: RD dan CI 95% di bawah himpunan "
               "penyesuaian M0–M6"); ng += 1


def ringkasan_kausal(doc, analisis, no):
    rows = []
    for n in analisis:
        est = _csv(C.OUT / "kausal" / n / "estimasi.csv", index_col=0)
        adj = NR._adj(C.OUT / "kausal" / n / "adjustment.txt")
        if est is None:
            continue
        e = est.estimasi
        rows.append({"analisis": n, "T → Y": f"{C.ANALISIS[n]['T']} → {C.ANALISIS[n]['Y']}",
                     "teridentifikasi": "ya" if adj and adj['Teridentifikasi'] == 'True' else "tidak (laten)",
                     "RD": e['aipw_RD'], "RD_lo": est.ci_bawah['aipw_RD'], "RD_hi": est.ci_atas['aipw_RD'],
                     "RR": e['aipw_RR'], "RR_lo": est.ci_bawah['aipw_RR'], "RR_hi": est.ci_atas['aipw_RR'],
                     "E-value": NR.evalue(e['aipw_RR'])})
    if rows:
        tabel(doc, gabung_ci(pd.DataFrame(rows)), f"Tabel 0.{no} Efek kausal AIPW (menu Inferensi Kausal)",
              {**F_RD, "E-value": "{:.2f}"})
        return no + 1
    return no


# ================================================================== ringkasan lintas analisis
def ringkasan_lintas(doc, bagian, no=1):
    import novelty
    bagian = [b for b in bagian if b in METODE_GRAF]
    if not bagian:
        return
    baris = []
    for n in C.ANALISIS:
        r = {"analisis": n, "T → Y": f"{C.ANALISIS[n]['T']} → {C.ANALISIS[n]['Y']}"}
        for b in bagian:
            f = (C.OUT / "baseline" / n / "rangkuman.csv") if b == "baseline" else \
                (C.OUT / "novelty" / b / n / "banding.csv")
            t = _csv(f)
            r[BAGIAN[b]] = t[t.terbaik].metode.iloc[0] if t is not None and t.terbaik.any() else "–"
        baris.append(r)
    tabel(doc, pd.DataFrame(baris), f"Tabel 0.{no} Metode terbaik per analisis pada setiap menu (kriteria menu "
          "masing-masing)")
    rows = []
    for b in [x for x in bagian if x != "baseline"]:
        for n in C.ANALISIS:
            t = _csv(C.OUT / "novelty" / b / n / "banding.csv")
            if t is None:
                continue
            k = novelty.METODE[b]['kode']
            x = t[t.kode == k]
            if len(x):
                x = x.iloc[0]
                rows.append({"kandidat": BAGIAN[b], "analisis": n, "RD": x.RD, "RD_lo": x.RD_lo, "RD_hi": x.RD_hi,
                             "detik_penemuan": x.get('detik_penemuan', np.nan), "memori_MB": x.get('memori_MB', np.nan),
                             "peringkat_rata2": x.peringkat_rata2,
                             "urutan": int(t.peringkat_rata2.rank(method='min')[t.kode == k].iloc[0])
                             if pd.notna(x.peringkat_rata2) else np.nan})
    if rows:
        tabel(doc, gabung_ci(pd.DataFrame(rows)), f"Tabel 0.{no + 1} Posisi setiap kandidat kebaruan terhadap baseline "
              "(urutan = posisi di antara metode yang diperingkat pada menu tersebut)", F_RD)


# ================================================================== dokumen
def buat(bagian=tuple(BAGIAN), analisis=None, simpan=None):
    """-> bytes .docx. bagian: subset BAGIAN; analisis: daftar nama analisis (None = semua)."""
    analisis = analisis or list(C.ANALISIS)
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)
    doc.add_heading(f"Laporan {BAGIAN[bagian[0]]}" if len(bagian) == 1 else
                    "Laporan Lengkap Analisis Kausal IDHS 2017" if len(bagian) == len(BAGIAN) else
                    "Laporan " + ", ".join(BAGIAN[b] for b in bagian), level=0)
    p = doc.add_paragraph()
    p.add_run("Inferensi kausal grafis pada data IDHS 2017 · dibuat otomatis dari aplikasi Kausal IDHS · "
              + time.strftime("%d-%m-%Y %H:%M")).italic = True
    doc.add_heading("Pengantar", level=1)
    graf = [b for b in bagian if b in METODE_GRAF]
    for t in [
        "Laporan ini menghimpun hasil, tabel, grafik, rangkuman proses, dan narasi dari menu "
        + ", ".join(BAGIAN[b] for b in bagian) + f" untuk {len(analisis)} analisis: "
        + "; ".join(f"{n} ({C.ANALISIS[n]['T']} → {C.ANALISIS[n]['Y']})" for n in analisis) + ".",
        "Seluruh estimasi memakai bobot sampel (v005/10⁶) dan desain survei berklaster (PSU v021, strata v023) "
        "IDHS 2017.",
    ] + ([
        "Semua metode memakai data kasus lengkap, tier waktu, dan langkah estimasi efek yang sama (adjustment set "
        "backdoor → g-komputasi berbobot survei → CI 95% bootstrap PSU di dalam strata), sehingga perbedaan hasil "
        "berasal dari cara graf diperoleh.",
        "Karena data nyata tidak memiliki ground truth, 'terbaik' berarti paling sesuai dengan kriteria menu "
        "(kecocokan dan presisi untuk baseline dan N1; waktu, memori, dan kecocokan/presisi untuk N2 dan N3), bukan "
        "jaminan paling benar secara kausal. Kandidat kebaruan adalah hipotesis desain yang belum teruji dan harus "
        "dibuktikan pada evaluasi semi-sintetik.",
    ] if graf else []):
        doc.add_paragraph(t).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if "kausal" in bagian or graf:
        doc.add_heading("Ringkasan lintas analisis", level=1)
        no = ringkasan_kausal(doc, analisis, 1) if "kausal" in bagian else 1
        ringkasan_lintas(doc, bagian, no)
    for bi, b in enumerate(bagian, start=1):
        doc.add_page_break()
        doc.add_heading(f"{bi}. {BAGIAN[b]}", level=1)
        if b in ("prep", "eda"):
            (isi_prep if b == "prep" else isi_eda)(doc, str(bi))
            continue
        if b == "benchmark_v4":                                          # simulasi: tidak per analisis
            isi_benchmark_v4(doc, str(bi))
            continue
        daftar = [n for n in analisis if b != "benchmark" or (C.OUT / "benchmark" / n / "putusan.csv").exists()]
        if not daftar:
            doc.add_paragraph("Hasil belum tersedia; jalankan menu ini di aplikasi terlebih dahulu.")
        for ai, n in enumerate(daftar, start=1):
            doc.add_heading(f"{bi}.{ai} {n}: {NR.lbl(C.ANALISIS[n]['T'])} → {NR.lbl(C.ANALISIS[n]['Y'])}", level=2)
            nomor = f"{bi}.{ai}"
            if b == "sol":
                isi_sol(doc, n, nomor)
            elif b == "benchmark":
                isi_benchmark(doc, n, nomor)
            elif b == "kausal":
                isi_kausal(doc, n, nomor)
            elif b == "baseline":
                isi_baseline(doc, n, nomor)
            else:
                isi_novelty(doc, b, n, nomor)
    buf = io.BytesIO(); doc.save(buf)
    if simpan:
        with open(simpan, "wb") as f:
            f.write(buf.getvalue())
    return buf.getvalue()
